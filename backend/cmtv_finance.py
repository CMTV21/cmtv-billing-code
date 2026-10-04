"""CMTV finances: one ledger for all income, costs and profit (CMTV local addition 2026-09-25).

Replaces the "CMTV Financials.xlsx" workflow (history imported once by scripts/import_finances.py):
- fin_transactions: every payment. source = "import" (the spreadsheet, up to 2026-09-30), "billing" (paid billing
  orders from the cut-over date, added automatically), or "manual" (recorded in Admin > Finances).
- fin_expenses: running costs by category.
- fin_config (one doc): cost per credit per server (with dates), credits per plan, PayPal fee, cut-over date.
The maths follows the spreadsheet: fee = round(amount*3.49% + 0.49, 2) for PayPal; cost = credits * cost per credit;
profit = amount - cost - fee; a month's costs = credit costs + PayPal fees + expenses (the Dashboard sheet's definition).
Rows are never hard-deleted: "deleted" rows are kept with deleted=True and ignored everywhere.
"""
import asyncio
import io
import logging
from datetime import datetime, timedelta
from typing import Optional

from bson import ObjectId
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/cmtv/finance", tags=["cmtv-finance"])
D = {}

METHODS = ["e-Transfer", "PayPal", "Other"]
EXPENSE_CATEGORIES = ["Hosting", "Telecom", "Software", "Platform", "Marketing", "Equipment", "Other"]
DEFAULT_CONFIG = {
    "_id": "config",
    "cutover": datetime(2026, 10, 1),
    "paypal_fee": {"pct": 0.0349, "fixed": 0.49},
    # cost per credit; a server can have several rates, each from a date (latest one on or before the payment date wins)
    "rates": [
        {"server": "CCTV", "cost_per_credit": 1.0, "from": None},
        {"server": "CCTV", "cost_per_credit": 1.25, "from": datetime(2026, 4, 1)},
        {"server": "Amethyst", "cost_per_credit": 3.5, "from": None},
        {"server": "OK4", "cost_per_credit": 6.0, "from": None},
        {"server": "Extreme", "cost_per_credit": 1.0, "from": None},
        {"server": "Stremio", "cost_per_credit": 0.0, "from": None},
        {"server": "VPN365", "cost_per_credit": 1.8, "from": None},
        {"server": "Imperium", "cost_per_credit": 3.0, "from": None},
        {"server": "Onetime", "cost_per_credit": 0.0, "from": None},
        {"server": "ABS", "cost_per_credit": 0.0, "from": None},
        {"server": "CMTVpn", "cost_per_credit": 0.0, "from": None},
        {"server": "CMTV+", "cost_per_credit": 0.0, "from": None},
    ],
    # credits a plan uses: (server, connections, months) -> credits. From the user's pricing sheet, 2026-09-25.
    "credits": (
        [{"server": "CCTV", "connections": c, "months": m, "credits": v} for c, row in {
            1: {1: 0.5, 3: 1.5, 6: 3, 12: 6}, 2: {1: 1, 3: 3, 6: 6, 12: 12},
            4: {1: 2, 3: 4, 6: 8, 12: 16}, 6: {1: 3, 3: 6, 6: 12, 12: 24}}.items() for m, v in row.items()]
        + [{"server": "Imperium", "connections": c, "months": m, "credits": v} for c, row in {
            1: {1: 1, 3: 3, 6: 6, 12: 12}, 2: {1: 2, 3: 4, 6: 8, 12: 14}, 3: {1: 3, 3: 6, 6: 10, 12: 16},
            4: {1: 4, 3: 8, 6: 12, 12: 18}, 5: {1: 5, 3: 10, 6: 14, 12: 20}}.items() for m, v in row.items()]
    ),
}


def init(**deps):
    D.update(deps)
    db = deps["db"]
    D["tx"], D["exp"], D["cfg"] = db.fin_transactions, db.fin_expenses, db.fin_config
    D["buys"], D["cfg_hist"] = db.fin_credit_purchases, db.fin_config_history   # 2026-09-27: credit purchases


async def startup():
    await D["tx"].create_index("date")
    await D["tx"].create_index("order_id", unique=True, partialFilterExpression={"order_id": {"$type": "string"}})
    await D["exp"].create_index("date")
    if not await D["cfg"].find_one({"_id": "config"}):
        await D["cfg"].insert_one(DEFAULT_CONFIG)
    asyncio.create_task(_sync_loop())


async def config():
    cfg = await D["cfg"].find_one({"_id": "config"}) or dict(DEFAULT_CONFIG)
    cfg["_avg"] = await avg_timelines()
    return cfg


# ---------------- credit purchases -> average cost per credit (2026-09-27) ----------------
# The user buys panel credits in batches at different prices (e.g. Imperium 100 at $3, then 200 at $2.50).
# Moving weighted average: each purchase is blended with the credits still on hand; sales (the ledger's credits)
# use credits up at the current average. From a server's first recorded purchase on, its cost per credit comes from
# this; before it, the dated rates below still apply, so older months don't change.

async def avg_timelines(server=None):
    """{server: {"points": [(date, avg)], "stock": credits on hand (estimate), "avg": current}} for servers with purchases"""
    q = {"deleted": {"$ne": True}}
    if server:
        q["server"] = server
    buys = await D["buys"].find(q).sort("date", 1).to_list(5000)
    out = {}
    for srv in sorted({b["server"] for b in buys}):
        mine = [b for b in buys if b["server"] == srv]
        first = mine[0]["date"]
        events = [(b["date"], 0, float(b["credits"]), float(b["total_paid"])) for b in mine]   # purchases first on a day
        async for t in D["tx"].find({"server": srv, "deleted": {"$ne": True}, "date": {"$gte": first}, "credits": {"$gt": 0}},
                                    {"date": 1, "credits": 1}):
            events.append((t["date"], 1, float(t.get("credits") or 0), 0.0))
        events.sort(key=lambda e: (e[0], e[1]))
        stock, avg, points = 0.0, None, []
        for when, kind, qty, paid in events:
            if kind == 0:
                avg = paid / qty if (avg is None or stock <= 0) else (stock * avg + paid) / (stock + qty)
                stock += qty
                points.append((when, round(avg, 4)))
            else:
                stock = max(0.0, stock - qty)
        out[srv] = {"points": points, "stock": round(stock, 2), "avg": round(avg or 0, 4), "first": first}
    return out


def cost_per_credit(cfg, server, when):
    tl = (cfg.get("_avg") or {}).get(server)
    if tl and tl["points"] and tl["points"][0][0] <= when:   # from purchases (average of credits on hand)
        best = tl["points"][0][1]
        for d, a in tl["points"]:
            if d <= when:
                best = a
        return best
    best, best_from = 0.0, None
    for r in cfg.get("rates", []):
        if r["server"].lower() != str(server).lower():
            continue
        f = r.get("from")
        if (f is None or f <= when) and (best_from is None or (f or datetime.min) >= best_from):
            best, best_from = float(r["cost_per_credit"]), (f or datetime.min)
    return best


def paypal_fee(cfg, method, amount):
    if method != "PayPal" or not amount:
        return 0.0
    p = cfg.get("paypal_fee") or {}
    return round(amount * float(p.get("pct", 0.0349)) + float(p.get("fixed", 0.49)), 2)


def complete(cfg, row):
    """Fill cost_per_credit, paypal_fee, total_cost and profit for a new/edited row"""
    amount = float(row.get("amount") or 0)
    credits = float(row.get("credits") or 0)
    cpc = row.get("cost_per_credit")
    cpc = cost_per_credit(cfg, row["server"], row["date"]) if cpc in (None, "") else float(cpc)
    fee = paypal_fee(cfg, row.get("method"), amount)
    total = round(credits * cpc, 2)
    row.update(amount=amount, credits=credits, cost_per_credit=cpc, paypal_fee=fee, total_cost=total,
               profit=round(amount - total - fee, 2))
    return row


def credits_for(cfg, server, connections, months):
    for c in cfg.get("credits", []):
        if c["server"] == server and int(c["connections"]) == int(connections or 0) and int(c["months"]) == int(months or 0):
            return float(c["credits"])
    return None


async def recost(server):
    """Re-work a server's credit cost and profit on ledger rows from its first recorded purchase on (the first time a row
    changes, its old figures are kept in cost_before_purchases). Returns how many rows changed."""
    cfg = await config()
    tl = cfg["_avg"].get(server)
    q = {"server": server, "deleted": {"$ne": True}, "credits": {"$gt": 0}}
    if tl:
        q["date"] = {"$gte": tl["first"]}
    else:   # no purchases left: rows that were re-costed go back to the dated rates
        q["cost_before_purchases"] = {"$exists": True}
    changed = 0
    async for t in D["tx"].find(q):
        row = {k: t.get(k) for k in ("server", "date", "amount", "credits", "method")}
        row["cost_per_credit"] = None
        complete(cfg, row)
        if abs(float(t.get("total_cost") or 0) - row["total_cost"]) < 0.005 and t.get("cost_per_credit") == row["cost_per_credit"]:
            continue
        upd = {"cost_per_credit": row["cost_per_credit"], "total_cost": row["total_cost"], "profit": row["profit"],
               "recosted_at": datetime.utcnow()}
        if "cost_before_purchases" not in t:
            upd["cost_before_purchases"] = {"cost_per_credit": t.get("cost_per_credit"), "total_cost": t.get("total_cost"),
                                            "profit": t.get("profit")}
        await D["tx"].update_one({"_id": t["_id"]}, {"$set": upd})
        changed += 1
    return changed


async def margins(cfg):
    """The pricing sheet: each plan in the credits table with its store price, credits, cost, profit and margin"""
    settings = await D["get_settings"]()
    groups = {g.get("id"): g.get("name") for g in settings.get("product_groups", [])}
    store = {}
    async for p in D["products"].find({"is_trial": {"$ne": True}, "account_type": {"$ne": "reseller"}}):
        prices = p.get("prices") or {}
        if not prices:
            continue
        months, price = next(iter(prices.items()))
        key = (_server_for(p, groups), int(p.get("max_connections") or 0), int(months))
        if float(price) > 0 and key not in store:
            store[key] = {"price": float(price), "product": (p.get("name") or "").strip()}
    now = datetime.utcnow()
    rows = []
    for c in sorted(cfg.get("credits", []), key=lambda c: (c["server"], int(c["connections"]), int(c["months"]))):
        key = (c["server"], int(c["connections"]), int(c["months"]))
        cpc = cost_per_credit(cfg, c["server"], now)
        s = store.get(key)
        cost = round(float(c["credits"]) * cpc, 2)
        rows.append({"server": c["server"], "connections": key[1], "months": key[2], "credits": float(c["credits"]),
                     "cost_per_credit": cpc, "cost": cost, "price": s["price"] if s else None, "product": s["product"] if s else None,
                     "profit": round(s["price"] - cost, 2) if s else None,
                     "margin": round((s["price"] - cost) * 100 / s["price"], 1) if s else None})
    return rows


# ---------------- billing orders -> ledger ----------------

def _server_for(product, group_names):
    """Ledger 'server' for a billing product (the spreadsheet's names)"""
    name = str((product or {}).get("name", "")).strip().lower()
    group = str(group_names.get((product or {}).get("group_id"), "")).lower()
    # CMTV local change 2026-10-01: Nuvio (retail + reseller credits) was filed as CCTV
    for key, server in (("audiobook", "ABS"), ("cmtvpn", "CMTVpn"), ("cmtv+", "CMTV+"), ("stremio", "Stremio"), ("nuvio", "Nuvio")):
        if key in name:
            return server
    if "imperium" in name or "imperium" in group or (product or {}).get("panel_type") == "aether":
        return "Imperium"
    return "CCTV"


def _method_for(order):
    m = str(order.get("payment_method") or "").lower()
    if m.startswith("paypal"):
        return "PayPal"
    if m in ("emt", "e-transfer", "etransfer", "interac"):
        return "e-Transfer"
    if m == "manual" and not order.get("payment_method_recorded"):
        return "e-Transfer"   # before 2026-09-25 every non-PayPal order said "manual"; nearly all were e-Transfers
    import cmtv_payments   # the recorded method's name (GhostPay, Manual, Cash, ...) - 2026-09-25
    return cmtv_payments.order_label(order)


async def sync_billing_orders():
    """Add paid billing orders from the cut-over date that aren't in the ledger yet. Safe to run any time."""
    cfg = await config()
    settings = await D["get_settings"]()
    groups = {g.get("id"): g.get("name") for g in settings.get("product_groups", [])}
    added, review = 0, 0
    async for o in D["orders"].find({"status": "paid", "paid_at": {"$gte": cfg["cutover"]}, "total": {"$gt": 0}}):
        oid = str(o["_id"])
        if await D["tx"].find_one({"order_id": oid}):
            continue
        user = await D["users"].find_one({"_id": ObjectId(o["user_id"])}) if ObjectId.is_valid(str(o.get("user_id"))) else None
        first_paid = await D["orders"].find_one({"user_id": o["user_id"], "status": "paid"}, sort=[("paid_at", 1)])
        credits, note = 0.0, []
        server = None
        for it in o.get("items", []):
            p = await D["products"].find_one({"_id": ObjectId(it["product_id"])}) if ObjectId.is_valid(str(it.get("product_id"))) else None
            server = server or _server_for(p, groups)
            months = int(it.get("term_months") or 1)
            if (p or {}).get("account_type") == "reseller":
                c = float(it.get("credits") or (p or {}).get("reseller_credits") or 0)   # 2026-09-28: chosen amount
            elif it.get("action_type") == "upgrade":   # 2026-10-04: device upgrade = the credits the panel actually charged
                c = float(it["panel_credits"]) if it.get("panel_credits") is not None else None
            else:
                c = credits_for(cfg, _server_for(p, groups), (p or {}).get("max_connections"), months)
            if c is None and cost_per_credit(cfg, _server_for(p, groups), o.get("paid_at") or datetime.utcnow()) == 0:
                c = 0.0   # own services (Stremio, CMTVpn, CMTV+, audiobooks) don't use paid credits
            if c is None:
                c = 0.0
                note.append(f"credits unknown for {it.get('product_name')}")
            credits += c
        row = complete(cfg, {
            "date": o.get("paid_at"), "server": server or "CCTV",
            "customer": (user or {}).get("name") or (user or {}).get("email") or "",
            "customer_email": (user or {}).get("email", ""), "user_id": o.get("user_id"),
            "new_user": bool(first_paid and first_paid["_id"] == o["_id"]),
            "amount": float(o.get("total") or 0), "method": _method_for(o), "credits": credits,
            "notes": "; ".join([i.get("product_name", "") for i in o.get("items", [])] + note),
            "source": "billing", "order_id": oid, "auto_renewal": bool(o.get("auto_renewal")),
            "needs_review": bool(note), "created_at": datetime.utcnow(),
        })
        try:
            await D["tx"].insert_one(row)
            added += 1
            review += bool(note)
        except Exception:   # duplicate order_id from a parallel run
            pass
    if added:
        logger.info(f"finance: {added} billing orders added to the ledger ({review} need a credits check)")
    return {"added": added, "needs_review": review}


async def _sync_loop():
    await asyncio.sleep(30)
    while True:
        try:
            await sync_billing_orders()
        except Exception as e:
            logger.error(f"finance sync failed: {type(e).__name__}: {e}")
        try:   # 2026-09-30: sales made outside billing -> Needs recording (cmtv_fin_inbox.py)
            import cmtv_fin_inbox
            await cmtv_fin_inbox.scan()
        except Exception as e:
            logger.error(f"finance inbox scan failed: {type(e).__name__}: {e}")
        await asyncio.sleep(600)


async def add_manual(data: dict, user: dict, extra: dict = None) -> dict:
    """A payment recorded by hand (Record payment, or a Needs recording item). Returns the saved row."""
    cfg = await config()
    server = str(data.get("server") or "").strip()
    if not server:
        raise HTTPException(status_code=400, detail="Choose a server")
    try:
        amount = float(data.get("amount"))
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Enter the amount received")
    row = complete(cfg, {
        "date": _clean_date(data.get("date")), "server": server, "customer": str(data.get("customer") or "").strip(),
        "new_user": bool(data.get("new_user")), "amount": amount,
        "method": data.get("method") if data.get("method") in METHODS else "Other",
        "credits": data.get("credits") or 0, "cost_per_credit": data.get("cost_per_credit"),
        "notes": str(data.get("notes") or "").strip(), "source": "manual",
        "created_at": datetime.utcnow(), "created_by": user.get("email"), **(extra or {}),
    })
    res = await D["tx"].insert_one(row)
    return _row_out(await D["tx"].find_one({"_id": res.inserted_id}))


# ---------------- summary ----------------

def _month_start(d):
    return d.replace(day=1, hour=0, minute=0, second=0, microsecond=0)


def _add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    return d.replace(year=d.year + y, month=m + 1, day=1)


async def _totals(start, end):
    live = {"date": {"$gte": start, "$lt": end}, "deleted": {"$ne": True}}
    t = await D["tx"].aggregate([{"$match": live}, {"$group": {
        "_id": None, "revenue": {"$sum": "$amount"}, "credit_cost": {"$sum": "$total_cost"},
        "fees": {"$sum": "$paypal_fee"}, "new_users": {"$sum": {"$cond": ["$new_user", 1, 0]}}, "count": {"$sum": 1}}}]).to_list(1)
    e = await D["exp"].aggregate([{"$match": live}, {"$group": {"_id": None, "total": {"$sum": "$amount"}}}]).to_list(1)
    t = t[0] if t else {"revenue": 0, "credit_cost": 0, "fees": 0, "new_users": 0, "count": 0}
    expenses = e[0]["total"] if e else 0.0
    costs = t["credit_cost"] + t["fees"] + expenses
    return {"revenue": round(t["revenue"], 2), "credit_cost": round(t["credit_cost"], 2), "fees": round(t["fees"], 2),
            "expenses": round(expenses, 2), "costs": round(costs, 2), "profit": round(t["revenue"] - costs, 2),
            "new_users": t["new_users"], "payments": t["count"]}


async def summary(month: Optional[str] = None):
    now = datetime.utcnow()
    m0 = datetime.strptime(month, "%Y-%m") if month else _month_start(now)
    m1 = _add_months(m0, 1)
    prev = _add_months(m0, -1)
    y0 = m0.replace(month=1)
    out = {
        "month": m0.strftime("%Y-%m"),
        "mtd": await _totals(m0, m1), "prev": await _totals(prev, m0), "ytd": await _totals(y0, m1),
        "all": await _totals(datetime(2000, 1, 1), datetime(2100, 1, 1)),
        "months": [],
    }
    for i in range(11, -1, -1):
        s = _add_months(m0, -i)
        tot = await _totals(s, _add_months(s, 1))
        out["months"].append({"month": s.strftime("%Y-%m"), **tot})
    by_server = await D["tx"].aggregate([{"$match": {"deleted": {"$ne": True}}}, {"$group": {
        "_id": "$server", "revenue": {"$sum": "$amount"}, "cost": {"$sum": "$total_cost"}, "fees": {"$sum": "$paypal_fee"},
        "count": {"$sum": 1}}}, {"$sort": {"revenue": -1}}]).to_list(100)
    out["servers"] = [{"server": s["_id"], "revenue": round(s["revenue"], 2), "cost": round(s["cost"] + s["fees"], 2),
                       "margin": round((s["revenue"] - s["cost"] - s["fees"]) / s["revenue"], 4) if s["revenue"] else None,
                       "payments": s["count"]} for s in by_server]
    cats = await D["exp"].aggregate([{"$match": {"deleted": {"$ne": True}}}, {"$group": {
        "_id": "$category", "total": {"$sum": "$amount"}}}, {"$sort": {"total": -1}}]).to_list(50)
    out["expense_categories"] = [{"category": c["_id"], "total": round(c["total"], 2)} for c in cats]
    # what only billing knows
    soon = now + timedelta(days=30)
    active = D["services"].find({"status": "active", "user_id": {"$nin": [None, ""]}}, {"expiry_date": 1, "panel_type": 1, "auto_renew": 1})
    n_active = due = 0
    async for s in active:
        n_active += 1
        e = s.get("expiry_date")
        if isinstance(e, str):
            try:
                e = datetime.fromisoformat(e.replace("Z", "").split(".")[0])
            except ValueError:
                e = None
        if isinstance(e, datetime) and now <= e <= soon:
            due += 1
    out["billing"] = {
        "active_services": n_active, "due_30d": due,
        "auto_renew": await D["services"].count_documents({"auto_renew.status": "ACTIVE"}),
        "paying_customers": len(await D["orders"].distinct("user_id", {"status": "paid"})),
        "needs_review": await D["tx"].count_documents({"needs_review": True, "deleted": {"$ne": True}}),
        "to_record": await D["db"].cmtv_fin_inbox.count_documents({"status": "open"}),   # 2026-09-30
    }
    return out


# ---------------- export ----------------

async def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Transaction Log"
    head = ["Date", "Server", "New User", "Amount Received", "Payment Method", "PayPal Fee", "Credits Used",
            "Cost per Credit", "Total Cost", "Total Profit", "Notes", "Source"]
    ws.append(head)
    async for t in D["tx"].find({"deleted": {"$ne": True}}).sort("date", 1):
        ws.append([t["date"], t.get("server"), "yes" if t.get("new_user") else "no", t.get("amount"), t.get("method") or "",
                   t.get("paypal_fee") or None, t.get("credits"), t.get("cost_per_credit"), t.get("total_cost"), t.get("profit"),
                   t.get("customer") if t.get("source") != "import" else (t.get("customer") or t.get("notes")), t.get("source")])
    e = wb.create_sheet("Expenses")
    e.append(["Date", "Category", "Description", "Amount", "Notes"])
    async for x in D["exp"].find({"deleted": {"$ne": True}}).sort("date", 1):
        e.append([x["date"], x.get("category"), x.get("description"), x.get("amount"), x.get("notes")])
    for sheet in (ws, e):
        for c in sheet[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F2A44")
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial")
                if isinstance(c.value, datetime):
                    c.number_format = "yyyy-mm-dd"
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------- API (admin only) ----------------

def _clean_date(v):
    try:
        return datetime.strptime(str(v)[:10], "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail="Enter the date as YYYY-MM-DD")


def _row_out(r):
    r["id"] = str(r.pop("_id"))
    return r


def init_routes():
    admin = D["get_current_admin_user"]

    @router.get("/summary")
    async def api_summary(month: Optional[str] = None, user=Depends(admin)):
        return await summary(month)

    @router.get("/config")
    async def api_config(user=Depends(admin)):
        c = await config()
        servers = sorted({r["server"] for r in c.get("rates", [])})
        return {"cutover": c["cutover"], "paypal_fee": c["paypal_fee"], "rates": c["rates"], "credits": c["credits"],
                "servers": servers, "methods": METHODS, "expense_categories": EXPENSE_CATEGORIES,
                # 2026-09-27: average cost per credit from purchases, as [date, avg] points per server
                "avg": {s: [[d, a] for d, a in v["points"]] for s, v in c["_avg"].items()}}

    @router.put("/config/rate")
    async def api_set_rate(data: dict, user=Depends(admin)):
        """Change a server's cost per credit from a date onward (earlier payments keep their old rate)"""
        server = str(data.get("server") or "").strip()
        if not server:
            raise HTTPException(status_code=400, detail="Server is required")
        rate = {"server": server, "cost_per_credit": float(data.get("cost_per_credit") or 0),
                "from": _clean_date(data["from"]) if data.get("from") else None}
        await D["cfg"].update_one({"_id": "config"}, {"$push": {"rates": rate}})
        return {"ok": True}

    # ---- credits & margins (2026-09-27) ----
    @router.get("/credits")
    async def api_credits(user=Depends(admin)):
        cfg = await config()
        now = datetime.utcnow()
        servers = sorted({c["server"] for c in cfg.get("credits", [])} | set(cfg["_avg"]))
        buys = await D["buys"].find({"deleted": {"$ne": True}}).sort("date", -1).to_list(500)
        return {
            "servers": [{"server": s, "cost_per_credit": cost_per_credit(cfg, s, now),
                         "from_purchases": s in cfg["_avg"], "stock": (cfg["_avg"].get(s) or {}).get("stock"),
                         "purchases": [{"id": str(b["_id"]), "date": b["date"], "credits": b["credits"], "total_paid": b["total_paid"],
                                        "per_credit": round(b["total_paid"] / b["credits"], 4) if b["credits"] else None,
                                        "notes": b.get("notes", "")} for b in buys if b["server"] == s]}
                        for s in servers],
            "all_servers": sorted({r["server"] for r in cfg.get("rates", [])} | set(servers)),
            "margins": await margins(cfg),
        }

    @router.post("/credits/purchases")
    async def api_add_purchase(data: dict, user=Depends(admin)):
        server = str(data.get("server") or "").strip()
        try:
            credits, paid = float(data.get("credits")), float(data.get("total_paid"))
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Enter the credits and what you paid")
        if not server or credits <= 0 or paid < 0:
            raise HTTPException(status_code=400, detail="Choose the server and enter credits above 0")
        await D["buys"].insert_one({"server": server, "date": _clean_date(data.get("date")), "credits": credits,
                                    "total_paid": round(paid, 2), "notes": str(data.get("notes") or "").strip()[:200],
                                    "created_at": datetime.utcnow(), "created_by": user.get("email")})
        changed = await recost(server)
        cfg = await config()
        return {"ok": True, "cost_per_credit": cost_per_credit(cfg, server, datetime.utcnow()), "recosted": changed}

    @router.delete("/credits/purchases/{buy_id}")
    async def api_remove_purchase(buy_id: str, user=Depends(admin)):
        b = await D["buys"].find_one({"_id": ObjectId(buy_id)}) if ObjectId.is_valid(buy_id) else None
        if not b:
            raise HTTPException(status_code=404, detail="Purchase not found")
        await D["buys"].update_one({"_id": b["_id"]}, {"$set": {"deleted": True, "deleted_at": datetime.utcnow(),
                                                                 "deleted_by": user.get("email")}})
        return {"ok": True, "recosted": await recost(b["server"])}

    @router.put("/credits/table")
    async def api_set_credits(data: dict, user=Depends(admin)):
        """Change (or add) how many credits a plan uses. The previous table is kept in fin_config_history."""
        try:
            server, conns, months = str(data["server"]).strip(), int(data["connections"]), int(data["months"])
            credits = float(data["credits"])
        except (KeyError, TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Server, devices, months and credits are needed")
        if not server or conns <= 0 or months <= 0 or credits < 0:
            raise HTTPException(status_code=400, detail="Devices and months must be above 0")
        cfg = await D["cfg"].find_one({"_id": "config"})
        table = list(cfg.get("credits", []))
        await D["cfg_hist"].insert_one({"credits": table, "saved_at": datetime.utcnow(), "by": user.get("email")})
        for c in table:
            if c["server"] == server and int(c["connections"]) == conns and int(c["months"]) == months:
                c["credits"] = credits
                break
        else:
            table.append({"server": server, "connections": conns, "months": months, "credits": credits})
        await D["cfg"].update_one({"_id": "config"}, {"$set": {"credits": table}})
        return {"ok": True}

    @router.get("/transactions")
    async def api_transactions(month: Optional[str] = None, q: str = "", source: str = "", user=Depends(admin)):
        f = {"deleted": {"$ne": True}}
        if month:
            m0 = datetime.strptime(month, "%Y-%m")
            f["date"] = {"$gte": m0, "$lt": _add_months(m0, 1)}
        if source:
            f["source"] = source
        if q:
            import re
            rx = {"$regex": re.escape(q.strip()[:80]), "$options": "i"}
            f["$or"] = [{"customer": rx}, {"notes": rx}, {"server": rx}]
        rows = await D["tx"].find(f).sort("date", -1).to_list(1000)
        return [_row_out(r) for r in rows]

    @router.post("/transactions")
    async def api_add_transaction(data: dict, user=Depends(admin)):
        return await add_manual(data, user)

    import cmtv_fin_inbox   # 2026-09-30: Needs recording (sales made outside billing)
    cmtv_fin_inbox.add_routes(router, admin)

    @router.put("/transactions/{tx_id}")
    async def api_edit_transaction(tx_id: str, data: dict, user=Depends(admin)):
        old = await D["tx"].find_one({"_id": ObjectId(tx_id)}) if ObjectId.is_valid(tx_id) else None
        if not old:
            raise HTTPException(status_code=404, detail="Payment not found")
        cfg = await config()
        row = {k: old[k] for k in ("date", "server", "customer", "new_user", "amount", "method", "credits", "notes") if k in old}
        for k in ("server", "customer", "notes", "method"):
            if k in data:
                row[k] = str(data[k] or "").strip()
        if "date" in data:
            row["date"] = _clean_date(data["date"])
        if "amount" in data:
            row["amount"] = float(data["amount"] or 0)
        if "credits" in data:
            row["credits"] = float(data["credits"] or 0)
        if "new_user" in data:
            row["new_user"] = bool(data["new_user"])
        row["cost_per_credit"] = data.get("cost_per_credit")
        complete(cfg, row)
        row.update(needs_review=False, updated_at=datetime.utcnow(), updated_by=user.get("email"))
        await D["tx"].update_one({"_id": old["_id"]}, {"$set": row})
        return _row_out(await D["tx"].find_one({"_id": old["_id"]}))

    @router.delete("/transactions/{tx_id}")
    async def api_remove_transaction(tx_id: str, user=Depends(admin)):
        res = await D["tx"].update_one({"_id": ObjectId(tx_id)} if ObjectId.is_valid(tx_id) else {"_id": None},
                                       {"$set": {"deleted": True, "deleted_at": datetime.utcnow(), "deleted_by": user.get("email")}})
        if not res.matched_count:
            raise HTTPException(status_code=404, detail="Payment not found")
        return {"ok": True}

    @router.get("/expenses")
    async def api_expenses(user=Depends(admin)):
        return [_row_out(r) for r in await D["exp"].find({"deleted": {"$ne": True}}).sort("date", -1).to_list(2000)]

    @router.post("/expenses")
    async def api_add_expense(data: dict, user=Depends(admin)):
        try:
            amount = float(data.get("amount"))
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Enter the amount")
        row = {"date": _clean_date(data.get("date")),
               "category": data.get("category") if data.get("category") in EXPENSE_CATEGORIES else "Other",
               "description": str(data.get("description") or "").strip(), "amount": round(amount, 2),
               "notes": str(data.get("notes") or "").strip(), "source": "manual",
               "created_at": datetime.utcnow(), "created_by": user.get("email")}
        res = await D["exp"].insert_one(row)
        return _row_out(await D["exp"].find_one({"_id": res.inserted_id}))

    @router.delete("/expenses/{exp_id}")
    async def api_remove_expense(exp_id: str, user=Depends(admin)):
        res = await D["exp"].update_one({"_id": ObjectId(exp_id)} if ObjectId.is_valid(exp_id) else {"_id": None},
                                        {"$set": {"deleted": True, "deleted_at": datetime.utcnow(), "deleted_by": user.get("email")}})
        if not res.matched_count:
            raise HTTPException(status_code=404, detail="Expense not found")
        return {"ok": True}

    @router.post("/sync")
    async def api_sync(user=Depends(admin)):
        return await sync_billing_orders()

    @router.get("/export")
    async def api_export(user=Depends(admin)):
        buf = await export_xlsx()
        name = f"CMTV-finances-{datetime.utcnow():%Y-%m-%d}.xlsx"
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f'attachment; filename="{name}"'})
